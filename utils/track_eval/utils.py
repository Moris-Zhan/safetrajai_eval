# utils.py
# -*- coding: utf-8 -*-
"""
基礎工具函式：型別轉換、MD5 指紋、影像尺寸、bbox 座標轉換、檔案探查等。
"""

import hashlib
import json
import re
from pathlib import Path

import cv2
import numpy as np


# =========================================================
# 型別 / 數值工具
# =========================================================
def _is_number(x):
    try:
        float(x)
        return True
    except Exception:
        return False


def _safe_float(x):
    try:
        if x is None:
            return None
        if hasattr(x, "item"):
            return float(x.item())
        return float(x)
    except Exception:
        return None


def _to_list(x):
    if x is None:
        return None
    if isinstance(x, np.ndarray):
        return x.tolist()
    if isinstance(x, (list, tuple)):
        return list(x)
    return [x]


# =========================================================
# Excel 安全值
# =========================================================
def _excel_safe(v, max_list_items=50):
    """openpyxl 無法直接寫入 list/np.ndarray/dict，需先轉換。"""
    if v is None:
        return None
    if isinstance(v, np.ndarray):
        v = v.tolist()
    if isinstance(v, (list, tuple)):
        if len(v) <= max_list_items:
            return str(v)
        return str(v[:max_list_items])[:-1] + f", ...] (len={len(v)})"
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def _ws_append_row(ws, row):
    """安全地將一列資料寫入 openpyxl worksheet。"""
    fixed = []
    for v in row:
        if isinstance(v, np.ndarray):
            v = v.tolist()
        if isinstance(v, (list, tuple, dict)):
            v = json.dumps(v, ensure_ascii=False)
        fixed.append(_safe_float(v) if _is_number(v) else v)
    ws.append(fixed)


def _autosize_columns(ws, max_width=60, scan_rows=300):
    """自動調整 worksheet 欄寬。"""
    max_rows_scan = min(ws.max_row, scan_rows)
    for col in range(1, ws.max_column + 1):
        best = 0
        for r in range(1, max_rows_scan + 1):
            v = ws.cell(r, col).value
            if v is None:
                continue
            best = max(best, len(str(v)))
        best = min(max_width, best + 2)
        from openpyxl.utils import get_column_letter
        ws.column_dimensions[get_column_letter(col)].width = max(10, best)


# =========================================================
# MD5 / 檔案指紋
# =========================================================
def file_md5(p: Path, chunk=1024 * 1024) -> str:
    h = hashlib.md5()
    with p.open("rb") as f:
        while True:
            b = f.read(chunk)
            if not b:
                break
            h.update(b)
    return h.hexdigest()


def md5_file(p: Path):
    b = p.read_bytes()
    return hashlib.md5(b).hexdigest(), len(b)


def labels_dir_fingerprint(labels_dir: str):
    """
    穩定的標籤目錄指紋：包含每個檔案的名稱、大小、內容 md5。
    """
    d = Path(labels_dir)
    files = sorted(d.glob("*.txt"))
    if not files:
        return ("<no_txt_files>", 0, 0, None, None)

    total_bytes = 0
    h = hashlib.md5()
    for p in files:
        b = p.read_bytes()
        total_bytes += len(b)
        h.update(p.name.encode("utf-8", errors="ignore"))
        h.update(str(len(b)).encode("utf-8"))
        h.update(hashlib.md5(b).hexdigest().encode("utf-8"))
    return (h.hexdigest(), len(files), total_bytes, files[0].name, files[-1].name)


# =========================================================
# 路徑 / 存在性檢查
# =========================================================
def _assert_exists(p: Path, what: str):
    if not p.exists():
        raise FileNotFoundError(f"{what} not found: {p}")


def ensure_parent(p: Path):
    """確保父目錄存在後回傳原 Path。"""
    p.parent.mkdir(parents=True, exist_ok=True)
    return p


# =========================================================
# 檔案探查（除錯用）
# =========================================================
def print_file_probe(title, p: Path, head_lines=3, tail_lines=3):
    p = Path(p)
    print(f"\n--- {title} ---")
    print("path:", p.resolve())
    if not p.exists():
        print("❌ NOT EXISTS")
        return
    h, n = md5_file(p)
    print("md5 :", h, "bytes:", n)
    lines = p.read_text(encoding="utf-8", errors="ignore").splitlines()
    print("lines:", len(lines))
    print("head:\n", "\n".join(lines[:head_lines]))
    if len(lines) > 0:
        print("tail:\n", "\n".join(lines[-tail_lines:]))


# =========================================================
# 影像 / 影片工具
# =========================================================
def get_video_wh(video_path: str):
    """回傳影片的 (width, height)。"""
    cap = cv2.VideoCapture(str(video_path))
    ok, frame = cap.read()
    cap.release()
    if not ok:
        raise RuntimeError(f"Cannot read video: {video_path}")
    h, w = frame.shape[:2]
    return w, h


# =========================================================
# YOLO bbox 座標轉換
# =========================================================
def yolo_xywhn_to_xywh_px(cx, cy, w, h, W, H):
    """YOLO 正規化中心點格式 → 像素左上角格式 (x1, y1, bw, bh)。"""
    cx, cy, w, h = float(cx), float(cy), float(w), float(h)
    bw = w * W
    bh = h * H
    x1 = (cx * W) - bw / 2
    y1 = (cy * H) - bh / 2
    return x1, y1, bw, bh


def clip_xywh(x, y, w, h, W, H):
    """將 bbox 裁切至影像邊界內。"""
    x2 = x + w
    y2 = y + h
    x = max(0.0, min(float(x), float(W - 1)))
    y = max(0.0, min(float(y), float(H - 1)))
    x2 = max(0.0, min(float(x2), float(W)))
    y2 = max(0.0, min(float(y2), float(H)))
    w = max(0.0, x2 - x)
    h = max(0.0, y2 - y)
    return x, y, w, h