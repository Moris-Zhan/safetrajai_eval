# excel_export.py
# -*- coding: utf-8 -*-
"""
TrackEval 多類別結果合併 + DET metrics 輸出到 Excel (.xlsx)。
"""

import json
import re
from pathlib import Path

import numpy as np
from openpyxl import Workbook

from evaluation.utils import (
    _safe_float, _to_list, _excel_safe,
    _ws_append_row, _autosize_columns,
)


# =========================================================
# TrackEval results flatten / merge -> Excel
# =========================================================
_key_re = re.compile(
    r"""
    ^(?:\[\d+\]\.)?
    (?P<dataset>[^.]+)\.
    (?P<tracker>[^.]+)\.
    (?P<seq>[^.]+)\.
    (?P<cls>[^.]+)\.
    (?P<family>HOTA|CLEAR|Identity|Count)\.
    (?P<metric>.+)
    $
    """,
    re.VERBOSE
)
_idx_re = re.compile(r"\[\d+\]")


def flatten_any(obj, prefix=""):
    """將任意巢狀 dict/list/ndarray 攤平為 (key, value) 列表。"""
    items = []
    if isinstance(obj, np.ndarray):
        obj = obj.tolist()
    if isinstance(obj, dict):
        for k, v in obj.items():
            p = f"{prefix}.{k}" if prefix else str(k)
            items.extend(flatten_any(v, p))
    elif isinstance(obj, (list, tuple)):
        for i, v in enumerate(obj):
            p = f"{prefix}[{i}]"
            items.extend(flatten_any(v, p))
    else:
        items.append((prefix, obj))
    return items


def _num_list_stats(v):
    """回傳 (mean, mid, min, max, n) 統計值。"""
    if v is None:
        return None, None, None, None, 0
    if isinstance(v, np.ndarray):
        v = v.tolist()
    if isinstance(v, (list, tuple)):
        arr = np.array(v, dtype=float).reshape(-1)
        n = int(arr.size)
        if n == 0:
            return None, None, None, None, 0
        mid = float(arr[n // 2])
        return float(arr.mean()), mid, float(arr.min()), float(arr.max()), n
    try:
        fv = float(v)
        return fv, fv, fv, fv, 1
    except Exception:
        return None, None, None, None, 0


def combine_trackeval_flat(flat_items):
    """
    將攤平後的 TrackEval 結果聚合成列表。

    Return:
      rows: list[dict]，每列包含 dataset/tracker/seq/class + <MetricKey>_mean/_mid/_min/_max/_n
      metric_keys: 排序後的 metric key 列表（如 "HOTA.HOTA", "CLEAR.sMOTA", …）
    """
    group = {}

    for k, v in flat_items:
        m = _key_re.match(k)
        if not m:
            continue
        d = m.groupdict()
        gkey = (d["dataset"], d["tracker"], d["seq"], d["cls"])

        metric_full = _idx_re.sub("", d["metric"])
        metric_leaf = metric_full.split(".")[-1]
        metric_key = f'{d["family"]}.{metric_leaf}'

        if isinstance(v, np.ndarray):
            v = v.tolist()

        group.setdefault(gkey, {}).setdefault(metric_key, [])
        if isinstance(v, (list, tuple)):
            if gkey[2] != "COMBINED_SEQ":
                group[gkey][metric_key].extend(v)
        else:
            if gkey[2] != "COMBINED_SEQ":
                group[gkey][metric_key].append(v)

    all_metric_keys = set()
    for md in group.values():
        all_metric_keys.update(md.keys())
    all_metric_keys = sorted(all_metric_keys)

    rows = []
    for (dataset, tracker, seq, cls), md in sorted(group.items()):
        row = {"dataset": dataset, "tracker": tracker, "seq": seq, "class": cls}
        if seq == "COMBINED_SEQ":
            continue
        for mk in all_metric_keys:
            vals = md.get(mk, None)
            if not vals:
                row[f"{mk}_mean"] = None
                row[f"{mk}_mid"] = None
                row[f"{mk}_min"] = None
                row[f"{mk}_max"] = None
                row[f"{mk}_n"] = None
                continue
            mean, mid, mn, mx, n = _num_list_stats(vals)
            row[f"{mk}_mean"] = mean
            row[f"{mk}_mid"] = mid
            row[f"{mk}_min"] = mn
            row[f"{mk}_max"] = mx
            row[f"{mk}_n"] = n
        rows.append(row)

    return rows, all_metric_keys


def sort_rows_grouped(rows, class_order=None):
    """
    依 dataset / tracker / class 分組排序，同組內依序排列。
    """
    if class_order is None:
        class_order = ["person", "car", "scooter", "small_truck", "bus", "truck"]
    class_rank = {c: i for i, c in enumerate(class_order)}

    def key(r):
        return (
            str(r.get("dataset", "")),
            str(r.get("tracker", "")),
            class_rank.get(str(r.get("class", "")), 999),
        )

    return sorted(rows, key=key)


def export_trackeval_results_multi_to_excel(
    all_results: dict,
    out_xlsx: str,
    classes_expected=None,
    tracker_base_hint="YOLO",
):
    """
    將多類別 TrackEval 結果合併輸出至 Excel。

    Returns:
        means_data: 各類別 mean 指標 list[dict]
        mid_data:   各類別 mid 指標 list[dict]
    """
    out_path = Path(out_xlsx)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    if classes_expected is None:
        classes_expected = ["person", "car", "scooter", "small_truck", "bus", "truck"]

    wb = Workbook()

    # Sheet 1: 所有原始 flat 資料
    ws_flat = wb.active
    ws_flat.title = "flat_all"
    _ws_append_row(ws_flat, ["run_class", "key_path", "value"])

    all_flat = []
    for cname, res in all_results.items():
        flat = flatten_any(res)
        for k, v in flat:
            _ws_append_row(ws_flat, [cname, k, _excel_safe(v)])
        all_flat.extend(flat)

    rows, metric_keys = combine_trackeval_flat(all_flat)

    # 補上 note 欄位
    for r in rows:
        r.setdefault("note", "")

    # 取得預設值（供補齊缺失類別用）
    if rows:
        any_row = rows[0]
        default_dataset = any_row.get("dataset", "MotChallenge2DBox")
        default_tracker = any_row.get("tracker", tracker_base_hint)
    else:
        default_dataset = "MotChallenge2DBox"
        default_tracker = tracker_base_hint

    # 確保每個預期類別至少出現一次
    present = set(r.get("class") for r in rows)
    for cls in classes_expected:
        if cls not in present:
            rows.append({
                "dataset": default_dataset,
                "tracker": f"{default_tracker}__{cls}",
                "class": cls,
                "note": "MISSING: no TrackEval result (likely missing TRK or skipped run)"
            })

    rows = sort_rows_grouped(rows, class_order=classes_expected)

    # Sheet 2: 所有欄位合併
    ws_comb = wb.create_sheet("combined_all")
    header = ["dataset", "tracker", "seq", "class", "note"]
    for mk in metric_keys:
        header += [f"{mk}_mean", f"{mk}_mid", f"{mk}_min", f"{mk}_max", f"{mk}_n"]
    _ws_append_row(ws_comb, header)
    for r in rows:
        _ws_append_row(ws_comb, [r.get(h) for h in header])

    # 關鍵指標清單（共用）
    key_metrics = [
        "HOTA.HOTA",
        "HOTA.DetA",
        "HOTA.AssA",
        "CLEAR.sMOTA",
        "Identity.IDF1",
        "Count.GT_Dets",
        "Count.Dets",
        "CLEAR.CLR_FP",
        "CLEAR.CLR_TP",
    ]

    # Sheet 3: key means
    ws_key = wb.create_sheet("key_means")
    key_header = ["dataset", "tracker", "seq", "class", "note"] + [f"{m}_mean" for m in key_metrics]
    _ws_append_row(ws_key, key_header)
    means_data = []
    for r in rows:
        line = [r.get("dataset"), r.get("tracker"), r.get("seq"), r.get("class"), r.get("note")]
        for m in key_metrics:
            line.append(r.get(f"{m}_mean"))
        _ws_append_row(ws_key, line)
        means_data.append(dict(zip(key_header, line)))

    # Sheet 4: key mid
    ws_mid = wb.create_sheet("key_mid")
    key_header_mid = ["dataset", "tracker", "seq", "class", "note"] + [f"{m}_mid" for m in key_metrics]
    _ws_append_row(ws_mid, key_header_mid)
    mid_data = []
    for r in rows:
        line = [r.get("dataset"), r.get("tracker"), r.get("seq"), r.get("class"), r.get("note")]
        for m in key_metrics:
            line.append(r.get(f"{m}_mid"))
        _ws_append_row(ws_mid, line)
        mid_data.append(dict(zip(key_header_mid, line)))

    for sheet in wb.worksheets:
        _autosize_columns(sheet)

    wb.save(out_path)
    print(f"✅ TrackEval multi-results saved: {out_path.resolve()}")
    return means_data, mid_data


# =========================================================
# DET metrics -> Excel
# =========================================================
def export_det_metrics_to_excel_openpyxl(
    metrics,
    names: dict,
    out_xlsx: str,
    include_curves=True,
    curve_max_points=400,
):
    """
    將 Ultralytics DET val metrics 輸出至 Excel。
    支援 summary、per-class AP、以及 PR/F1 曲線（可選）。
    """
    out_path = Path(out_xlsx)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws_summary = wb.active
    ws_summary.title = "summary"

    box = getattr(metrics, "box", None)
    if box is None:
        raise RuntimeError("metrics.box is None")

    _ws_append_row(ws_summary, ["mAP@50:95", "mAP@50", "mAP@75"])
    _ws_append_row(ws_summary, [
        _safe_float(getattr(box, "map", None)),
        _safe_float(getattr(box, "map50", None)),
        _safe_float(getattr(box, "map75", None)),
    ])

    # 每類別 AP
    ws_pc = wb.create_sheet("per_class_AP")
    maps = getattr(box, "maps", None)
    all_ap = getattr(box, "all_ap", None)
    ap_ci = getattr(box, "ap_class_index", None)

    cls2row = None
    if ap_ci is not None:
        ap_ci_list = [int(x) for x in _to_list(ap_ci)]
        cls2row = {cid: i for i, cid in enumerate(ap_ci_list)}

    _ws_append_row(ws_pc, ["class_id", "class_name", "mAP@50:95", "AP@0.50", "AP@0.75", "AP@0.95"])
    maps_list = _to_list(maps) if maps is not None else []

    for cid in sorted(names.keys()):
        cname = names.get(cid, str(cid))
        map5095 = _safe_float(maps_list[cid]) if maps_list and cid < len(maps_list) else None

        ap50 = ap75 = ap95 = None
        if all_ap is not None and cls2row is not None and cid in cls2row:
            r = cls2row[cid]
            row = _to_list(all_ap)[r]
            if isinstance(row, (list, tuple)) and len(row) >= 10:
                ap50 = _safe_float(row[0])
                ap75 = _safe_float(row[5])
                ap95 = _safe_float(row[9])
            elif isinstance(row, (list, tuple)) and len(row) >= 1:
                ap50 = _safe_float(row[0])

        _ws_append_row(ws_pc, [cid, cname, map5095, ap50, ap75, ap95])

    # 曲線資料（可選）
    if include_curves:
        curves_results = getattr(box, "curves_results", None)
        if curves_results is not None:
            crs = _to_list(curves_results)
            for idx, item in enumerate(crs):
                if not isinstance(item, (list, tuple)):
                    item = [item]

                x_label = y_label = ""
                if len(item) >= 3 and isinstance(item[-1], str) and isinstance(item[-2], str):
                    x_label, y_label = item[-2], item[-1]

                data = [x for x in item if not isinstance(x, str)]
                ws = wb.create_sheet(f"curve_{idx}"[:31])
                _ws_append_row(ws, ["x_label", x_label])
                _ws_append_row(ws, ["y_label", y_label])
                _ws_append_row(ws, [])

                def downsample_idx(n, maxn):
                    if n <= maxn:
                        return list(range(n))
                    step = max(1, n // maxn)
                    return list(range(0, n, step))

                if len(data) == 1:
                    y = _to_list(data[0])
                    if y and isinstance(y[0], list):
                        N = len(y[0])
                        pick = downsample_idx(N, curve_max_points)
                        _ws_append_row(ws, ["class_id", "class_name"] + [f"p_{j}" for j in range(len(pick))])
                        for cid in range(len(y)):
                            row = y[cid]
                            _ws_append_row(ws, [cid, names.get(cid, str(cid))] + [row[i] for i in pick])

                elif len(data) >= 2:
                    x = _to_list(data[0])
                    y = _to_list(data[1])
                    pick = downsample_idx(len(x), curve_max_points) if x is not None else None
                    if x is not None:
                        _ws_append_row(ws, ["x_points"] + [x[i] for i in pick])

                    if y and isinstance(y[0], list):
                        N = len(y[0])
                        if pick is None:
                            pick = downsample_idx(N, curve_max_points)
                        _ws_append_row(ws, ["class_id", "class_name"] + [f"y_{j}" for j in range(len(pick))])
                        for cid in range(len(y)):
                            row = y[cid]
                            _ws_append_row(ws, [cid, names.get(cid, str(cid))] + [row[i] for i in pick])

    for ws in wb.worksheets:
        _autosize_columns(ws)

    wb.save(out_path)
    print(f"✅ DET metrics saved: {out_path.resolve()}")


# =========================================================
# 終端評估摘要（彩色輸出）
# =========================================================
# def print_eval_summary(means_data, mid_data, seq_name=None):
#     """在終端印出關鍵指標的 mean / mid 摘要表格。"""
#     GREEN = "\033[32m"
#     RESET = "\033[0m"
#     key_metrics = [
#         "HOTA.HOTA",
#         "HOTA.DetA",
#         "HOTA.AssA",
#         # "CLEAR.sMOTA",
#         # "Identity.IDF1",
#         # "Count.GT_Dets",
#         # "Count.Dets",
#         # "CLEAR.CLR_FP",
#         # "CLEAR.CLR_TP",
#     ]

#     # mid_data 轉成 {class: row} 方便對照
#     mid_lookup = {r.get("class"): r for r in mid_data}

#     title = f"序列：{seq_name}" if seq_name else "評估結果摘要"
#     print(GREEN + "\n" + "=" * 90 + RESET)
#     print(GREEN + f"  {title}" + RESET)
#     print(GREEN + "=" * 90 + RESET)

#     # 表頭
#     col_w = 14
#     metric_labels = [m.split(".")[-1] for m in key_metrics]   # 只取最後一段，較簡短
#     header_parts = f"{'類別':<14}" + "".join(f"{'[mean]'+lb:>{col_w}}{'[mid]'+lb:>{col_w}}" for lb in metric_labels)
#     print(GREEN + header_parts + RESET)
#     print(GREEN + "-" * len(header_parts) + RESET)

#     for row in means_data:
#         cls = row.get("class", "?")
#         note = row.get("note", "")
#         mid_row = mid_lookup.get(cls, {})

#         if note:
#             print(GREEN + f"{cls:<14}  ⚠️  {note}" + RESET)
#             continue

#         values = ""
#         for m in key_metrics:
#             mean_v = row.get(f"{m}_mean")
#             mid_v  = mid_row.get(f"{m}_mid")
#             mean_s = f"{mean_v:>{col_w}.4f}" if mean_v is not None else f"{'N/A':>{col_w}}"
#             mid_s  = f"{mid_v:>{col_w}.4f}" if mid_v  is not None else f"{'N/A':>{col_w}}"
#             values += mean_s + mid_s

#         print(GREEN + f"{cls:<14}{values}" + RESET)

#     print(GREEN + "=" * 90 + RESET)


def print_eval_summary(means_data, mid_data, seq_name=None):
    """在終端印出 '一列式' mean/mid 摘要表格，mean 和 mid 各指標依據需求排序。"""
    GREEN = "\033[32m"
    RESET = "\033[0m"

    # 按指定順序：mean-HOTA, mean-DetA, mean-AssA, mid-HOTA, mid-DetA, mid-AssA
    # key_metrics 只做對應欄位查詢
    mean_metrics = [
        "HOTA.HOTA",
        "HOTA.DetA",
        "HOTA.AssA"
    ]
    mid_metrics = [
        "HOTA.HOTA",
        "HOTA.DetA",
        "HOTA.AssA"
    ]

    # mid_data 轉成 {class: row} 方便對照
    mid_lookup = {r.get("class"): r for r in mid_data}
    title = f"序列：{seq_name}" if seq_name else "評估結果摘要"
    print(GREEN + "\n" + "=" * 90 + RESET)
    print(GREEN + f"  {title}" + RESET)
    print(GREEN + "=" * 90 + RESET)

    col_w = 14

    # 新表頭
    header_parts = f"{'類別':<14}"
    for m in mean_metrics:
        lb = m.split('.')[-1]
        header_parts += f"{('[mean]'+lb):>{col_w}}"
    for m in mid_metrics:
        lb = m.split('.')[-1]
        header_parts += f"{('[mid]'+lb):>{col_w}}"
    print(GREEN + header_parts + RESET)
    print(GREEN + "-" * len(header_parts) + RESET)

    for row in means_data:
        cls = row.get("class", "?")
        note = row.get("note", "")
        mid_row = mid_lookup.get(cls, {})
        if note:
            print(GREEN + f"{cls:<14}  ⚠️  {note}" + RESET)
            continue
        values = ""
        for m in mean_metrics:
            mean_v = row.get(f"{m}_mean")
            mean_s = f"{mean_v:>{col_w}.9f}" if mean_v is not None else f"{'N/A':>{col_w}}"
            values += mean_s
        for m in mid_metrics:
            mid_v = mid_row.get(f"{m}_mid")
            mid_s = f"{mid_v:>{col_w}.9f}" if mid_v is not None else f"{'N/A':>{col_w}}"
            values += mid_s
        print(GREEN + f"{cls:<14}{values}" + RESET)
    print(GREEN + "=" * 90 + RESET)
