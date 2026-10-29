"""Default class lists in the evaluation code follow the class ID contract."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

import safetrajai.evaluation.excel_export as excel_export
import safetrajai.evaluation.trackeval_runner as trackeval_runner
from safetrajai.classes import CLASS_NAMES

CONTRACT_ORDER = ["person", "car", "scooter", "small_truck", "bus", "truck"]
ALPHABETICAL_ORDER = sorted(CONTRACT_ORDER)


def test_contract_order_is_the_class_id_order() -> None:
    assert list(CLASS_NAMES.values()) == CONTRACT_ORDER


def test_export_fills_every_class_by_default(tmp_path: Path) -> None:
    means_data, mid_data = excel_export.export_trackeval_results_multi_to_excel({}, str(tmp_path / "out.xlsx"))

    # Filled rows carry a per-class tracker name, so they sort by tracker (alphabetically), not by class rank.
    assert [row["class"] for row in means_data] == ALPHABETICAL_ORDER
    assert [row["class"] for row in mid_data] == ALPHABETICAL_ORDER
    assert all(row["note"].startswith("MISSING") for row in means_data)

    sheet = load_workbook(tmp_path / "out.xlsx")["key_means"]
    assert [row[3] for row in sheet.iter_rows(min_row=2, values_only=True)] == ALPHABETICAL_ORDER


def test_export_honours_an_explicit_class_list(tmp_path: Path) -> None:
    means_data, _ = excel_export.export_trackeval_results_multi_to_excel(
        {}, str(tmp_path / "out.xlsx"), classes_expected=["truck", "car"]
    )
    assert [row["class"] for row in means_data] == ["car", "truck"]


def test_sort_rows_grouped_defaults_to_contract_order() -> None:
    rows = [{"class": name} for name in reversed(CONTRACT_ORDER)]
    assert [row["class"] for row in excel_export.sort_rows_grouped(rows)] == CONTRACT_ORDER


def test_sort_rows_grouped_puts_unknown_classes_last() -> None:
    rows = [{"class": "tram"}, {"class": "car"}, {"class": "person"}]
    assert [row["class"] for row in excel_export.sort_rows_grouped(rows)] == ["person", "car", "tram"]


@pytest.mark.parametrize("module", [excel_export, trackeval_runner])
def test_modules_do_not_hardcode_the_class_list(module) -> None:
    source = Path(module.__file__).read_text(encoding="utf-8")
    assert '"person", "car", "scooter", "small_truck", "bus", "truck"' not in source
