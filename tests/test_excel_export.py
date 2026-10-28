"""Cell-by-cell snapshots of the Excel reports written by evaluation.excel_export."""

import json
from pathlib import Path
from types import SimpleNamespace

import numpy as np
from openpyxl import load_workbook

import safetrajai.evaluation.excel_export as excel_export

SNAPSHOT_DIR = Path(__file__).parent / "fixtures" / "excel_export"
DATASET = "MotChallenge2DBox"
ALPHAS = [round(0.05 * i, 2) for i in range(1, 20)]


def _class_metrics(scale: float, with_counts: bool = True) -> dict:
    metrics = {
        "HOTA": {
            "HOTA": np.array([scale * a for a in ALPHAS]),
            "DetA": [scale * a / 2 for a in ALPHAS],
            "AssA": [scale * (1 - a) for a in ALPHAS],
        },
        "CLEAR": {"sMOTA": scale * 0.5, "CLR_FP": int(scale * 100), "CLR_TP": int(scale * 200)},
        "Identity": {"IDF1": scale * 0.25},
    }
    if with_counts:
        metrics["Count"] = {"GT_Dets": int(scale * 300), "Dets": int(scale * 280)}
    return metrics


def synthetic_trackeval_results() -> dict:
    """Per-class TrackEval results: two sequences, a COMBINED_SEQ row, one class without counts, two classes absent."""
    results = {}
    for scale, name, with_counts in ((0.9, "person", True), (0.6, "car", True), (0.3, "bus", False), (0.1, "scooter", True)):
        tracker = f"YOLO11x__{name}"
        results[name] = {DATASET: {tracker: {
            "seqA": {name: _class_metrics(scale, with_counts)},
            "seqB": {name: _class_metrics(scale / 2, with_counts)},
            "COMBINED_SEQ": {name: _class_metrics(scale / 3, with_counts)},
        }}}
    return results


def dump_workbook(path: Path) -> dict[str, list[list]]:
    workbook = load_workbook(path)
    return {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)] for sheet in workbook.worksheets}


def export_trackeval_snapshot(tmp_dir: Path) -> dict:
    out = tmp_dir / "trackeval.xlsx"
    means_data, mid_data = excel_export.export_trackeval_results_multi_to_excel(
        synthetic_trackeval_results(), str(out), tracker_base_hint="YOLO11x"
    )
    sheets = dump_workbook(out)
    flat_rows = len(sheets.pop("flat_all"))  # raw dump; only its size is pinned to keep the fixture small
    return {"flat_all_rows": flat_rows, "sheets": sheets, "means_data": means_data, "mid_data": mid_data}


def test_trackeval_workbook_matches_snapshot(tmp_path: Path) -> None:
    expected = json.loads((SNAPSHOT_DIR / "trackeval_multi.json").read_text(encoding="utf-8"))
    actual = json.loads(json.dumps(export_trackeval_snapshot(tmp_path)))
    assert actual == expected


def test_trackeval_snapshot_covers_the_key_sheets(tmp_path: Path) -> None:
    sheets = export_trackeval_snapshot(tmp_path)["sheets"]
    assert {"key_means", "key_mid"} <= set(sheets)
    header = sheets["key_means"][0]
    count_column = header.index("Count.GT_Dets_mean")
    assert any(row[3] == "bus" and row[count_column] is None for row in sheets["key_means"][1:])
    assert any(row[4] and row[4].startswith("MISSING") for row in sheets["key_mid"][1:])


def synthetic_det_metrics() -> SimpleNamespace:
    """Ultralytics-like val metrics whose curves hit every branch of the DET curve writer."""
    ramp = lambda n, scale: [round(scale * i / n, 4) for i in range(n)]  # noqa: E731
    per_class = lambda rows, n: [ramp(n, 1.0 + k) for k in range(rows)]  # noqa: E731
    curves = [
        [ramp(25, 1.0), per_class(2, 25), "Recall", "Precision"],  # x + per-class y, 25 points
        [per_class(3, 8), "Confidence", "F1"],  # y only, 3 rows (class 2 has no name), 8 points
        [ramp(10, 1.0), ramp(10, 2.0)],  # x + flat y: only x_points is written
        [ramp(15, 1.0), per_class(2, 15), "Confidence", "Precision"],  # 15 points
        np.array(ramp(6, 3.0)),  # single flat array: labels only
        [None, per_class(2, 5)],  # no x: y rows are down-sampled on their own
    ]
    box = SimpleNamespace(
        map=0.5, map50=0.75, map75=0.55,
        maps=np.array([0.5, 0.4]),
        all_ap=np.array([ramp(10, 1.0), ramp(10, 0.5)]),
        ap_class_index=np.array([0, 1]),
        curves_results=curves,
    )
    return SimpleNamespace(box=box)


DET_NAMES = {0: "person", 1: "car", 5: "truck"}
DET_VARIANTS = {
    "curves_max_10": {"include_curves": True, "curve_max_points": 10},
    "curves_default": {"include_curves": True},
    "no_curves": {"include_curves": False},
}


def export_det_snapshot(tmp_dir: Path) -> dict:
    snapshot = {}
    for name, options in DET_VARIANTS.items():
        out = tmp_dir / f"{name}.xlsx"
        excel_export.export_det_metrics_to_excel_openpyxl(synthetic_det_metrics(), DET_NAMES, str(out), **options)
        snapshot[name] = dump_workbook(out)
    return snapshot


def test_det_workbook_matches_snapshot(tmp_path: Path) -> None:
    expected = json.loads((SNAPSHOT_DIR / "det_metrics.json").read_text(encoding="utf-8"))
    actual = json.loads(json.dumps(export_det_snapshot(tmp_path)))
    assert actual == expected


def test_det_snapshot_keeps_the_oversampling_of_long_curves(tmp_path: Path) -> None:
    sheet = export_det_snapshot(tmp_path)["curves_max_10"]["curve_0"]
    x_points = next(row for row in sheet if row[0] == "x_points")
    assert len([value for value in x_points[1:] if value is not None]) == 13  # 25 points, step 25 // 10 = 2
